"""
Script to generate TestCase_PayOrder.docx from test files and auto-populate
JUnit test results by parsing the Surefire XML reports.
Run: python generate_testcase_doc.py
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
import xml.etree.ElementTree as ET

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "TestCase_PayOrder.docx")
REPORTS_DIR = os.path.join(os.path.dirname(__file__), "..", "aims-backend", "target", "surefire-reports")

# ─────────────────────────────────────────────
# Test case mapping for auto-populating results
# ─────────────────────────────────────────────
TEST_CASES = [
    # ── PayOrderController – generate-qr ──
    {
        "id": "UT001",
        "name": "Generate QR Code Successfully",
        "method": "shouldGenerateQRCodeSuccessfully",
        "description": "Successfully generate a payment QR code when input data is valid and VietQR returns a valid QR code.",
        "unit": "PayOrderController\n.generateQRCode()",
        "input": "POST /api/payment/generate-qr\n{\n  invoiceId: 1,\n  shippingFee: 10000,\n  totalProductPriceExVAT: 90000,\n  totalProductPriceIncVAT: 100000,\n  totalAmount: 110000\n}",
        "expected": "HTTP 200 OK\n{\n  success: true,\n  message: \"QR code generated\",\n  data.qrCode.qrCode: \"000201...\"\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Happy path – valid input",
    },
    {
        "id": "UT002",
        "name": "Handle InvalidTokenException When Generating QR",
        "method": "shouldHandleInvalidTokenExceptionWhenGeneratingQR",
        "description": "Handle error when VietQR cannot authenticate the access token (token expired or invalid).",
        "unit": "PayOrderController\n.generateQRCode()",
        "input": "POST /api/payment/generate-qr\n(VietQR API throws InvalidTokenException:\n \"Token expired\")",
        "expected": "HTTP 503 Service Unavailable\n{\n  success: false,\n  message: \"Service unavailable:\n  Token expired\"\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Alt flow:\nAuthentication Failure",
    },
    {
        "id": "UT003",
        "name": "Handle QRCodeGenerationException",
        "method": "shouldHandleQRCodeGenerationException",
        "description": "Handle error when the VietQR API fails to generate a QR code (e.g., Bank API is down).",
        "unit": "PayOrderController\n.generateQRCode()",
        "input": "POST /api/payment/generate-qr\n(VietQR API throws QRCodeGenerationException:\n \"Bank API down\")",
        "expected": "HTTP 502 Bad Gateway\n{\n  success: false\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Alt flow:\nQR Generation Failure",
    },
    {
        "id": "UT004",
        "name": "Handle Unknown Exception When Generating QR",
        "method": "shouldHandleUnknownExceptionWhenGeneratingQR",
        "description": "Handle unexpected runtime errors during QR generation (e.g., database down, unhandled exception).",
        "unit": "PayOrderController\n.generateQRCode()",
        "input": "POST /api/payment/generate-qr\n(throws RuntimeException:\n \"Database down\")",
        "expected": "HTTP 500 Internal Server Error\n{\n  success: false\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Alt flow:\nUnknown Error",
    },
    # ── PayOrderController – confirm ──
    {
        "id": "UT005",
        "name": "Confirm Payment – Cache Hit",
        "method": "shouldConfirmPaymentImmediatelyIfCacheHit",
        "description": "Payment confirmation returns SUCCESS immediately when the transaction result already exists in cache (webhook was processed earlier).",
        "unit": "PayOrderController\n.confirmPayment()",
        "input": "POST /api/payment/confirm\n{\n  paymentRef: \"REF-TEST\"\n}\n(Cache already has TransactionInfo\n transactionId=123)",
        "expected": "HTTP 200 OK\n{\n  data.status: \"SUCCESS\"\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Cache hit – no need to re-call VietQR",
    },
    {
        "id": "UT006",
        "name": "Confirm Payment – Webhook Slow (Pending)",
        "method": "shouldReturnPendingIfWebhookSlow",
        "description": "Returns PENDING status when the VietQR webhook has not yet called back within the 700ms wait window.",
        "unit": "PayOrderController\n.confirmPayment()",
        "input": "POST /api/payment/confirm\n{\n  paymentRef: \"REF-TEST\"\n}\n(Cache: completedPayment=null,\n pendingInvoice=Invoice{})",
        "expected": "HTTP 200 OK\n{\n  data.status: \"PENDING\"\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Webhook not received within 700ms",
    },
    {
        "id": "UT007",
        "name": "Confirm Payment – Session Not Found",
        "method": "shouldThrowUnknownExceptionIfSessionNotFound",
        "description": "Throws an error when no payment session is found for the given paymentRef (not in cache or already expired).",
        "unit": "PayOrderController\n.confirmPayment()",
        "input": "POST /api/payment/confirm\n{\n  paymentRef: \"REF-TEST\"\n}\n(Cache: completedPayment=null,\n pendingInvoice=null)",
        "expected": "HTTP 500 Internal Server Error\n(throws UnknownException:\n \"Payment session not found\n  or expired\")",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Session does not exist or has expired",
    },
    {
        "id": "UT008",
        "name": "Confirm Payment – Timeout",
        "method": "shouldHandlePaymentTimeoutException",
        "description": "Handles payment timeout scenario when PaymentTimeoutException is thrown.",
        "unit": "PayOrderController\n.confirmPayment()",
        "input": "POST /api/payment/confirm\n{\n  paymentRef: \"REF-TEST\"\n}\n(throws PaymentTimeoutException)",
        "expected": "HTTP 408 Request Timeout",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Alt flow:\nPayment Timeout",
    },
    {
        "id": "UT009",
        "name": "Confirm Payment – Callback Validation Fails",
        "method": "shouldHandleCallbackValidationException",
        "description": "Handles tampered or invalid callback data when CallbackValidationException is thrown.",
        "unit": "PayOrderController\n.confirmPayment()",
        "input": "POST /api/payment/confirm\n{\n  paymentRef: \"REF-TEST\"\n}\n(throws CallbackValidationException:\n \"Tampered\")",
        "expected": "HTTP 400 Bad Request\n{\n  message: \"Security verification\n  failed: Tampered\"\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Alt flow:\nData Tampered",
    },
    # ── PayOrderController – switch-method ──
    {
        "id": "UT010",
        "name": "Switch Payment Method to PayPal",
        "method": "shouldSwitchMethodToPayPal",
        "description": "Successfully switches the payment method to PayPal.",
        "unit": "PayOrderController\n.switchMethod()",
        "input": "POST /api/payment/switch-method\n{\n  paymentRef: \"REF-TEST\",\n  method: \"PayPal\"\n}",
        "expected": "HTTP 200 OK\n{\n  data.method: \"PayPal\"\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Switch to PayPal method",
    },
    # ── PayOrderController – transaction/{ref} ──
    {
        "id": "UT011",
        "name": "Get Transaction Info – Found",
        "method": "shouldReturnTransactionIfFound",
        "description": "Successfully retrieves transaction information when the transaction already exists in cache.",
        "unit": "PayOrderController\n.returnTransactionInfo()",
        "input": "GET /api/payment/transaction/REF-TEST\n(Cache has TransactionInfo\n amount=500)",
        "expected": "HTTP 200 OK\n{\n  data.amount: 500.0\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Transaction already completed",
    },
    # ── VietQRInboundController – token_generate ──
    {
        "id": "UT012",
        "name": "Get Token – Valid Basic Auth",
        "method": "shouldReturnTokenWhenBasicAuthIsValid",
        "description": "VietQR successfully retrieves the Bearer token when sending valid Basic Auth credentials (admin:admin).",
        "unit": "VietQRInboundController\n.getToken()",
        "input": "POST /vqr/api/token_generate\nAuthorization: Basic YWRtaW46YWRtaW4=\n(admin:admin Base64 encoded)",
        "expected": "HTTP 200 OK\n{\n  access_token: \"dummy_token_123\"\n}",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Valid Basic Auth credentials",
    },
    {
        "id": "UT013",
        "name": "Get Token – Invalid Basic Auth",
        "method": "shouldReturnUnauthorizedWhenBasicAuthIsInvalid",
        "description": "Returns 401 Unauthorized when the Basic Auth credentials in the request are incorrect.",
        "unit": "VietQRInboundController\n.getToken()",
        "input": "POST /vqr/api/token_generate\nAuthorization: Basic WRONG",
        "expected": "HTTP 401 Unauthorized",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Invalid Basic Auth credentials",
    },
    # ── VietQRInboundController – callback ──
    {
        "id": "UT014",
        "name": "Process Callback Successfully",
        "method": "shouldProcessCallbackSuccessfully",
        "description": "Successfully processes a VietQR payment webhook: saves Invoice, creates and saves TransactionInfo, and updates the cache.",
        "unit": "VietQRInboundController\n.receiveCallback()",
        "input": "POST /vqr/bank/api/test/\ntransaction-callback\nAuthorization: Bearer dummy_token_123\n{\n  status: \"00\",\n  amount: 100000,\n  content: \"Payment AIMS REF-12345\",\n  transactionRefId: \"REF-12345\"\n}\n(Cache has Invoice for REF-12345)",
        "expected": "HTTP 200 OK\n{\n  success: true\n}\ninvoiceRepository.save() called 2x\ntransactionRepository.save() called 1x\nCache.addCompletedPayment() called 1x",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Main success flow",
    },
    {
        "id": "UT015",
        "name": "Callback – Invalid Bearer Token",
        "method": "shouldThrowCallbackValidationExceptionWhenBearerIsInvalid",
        "description": "Throws CallbackValidationException when the Bearer token in the webhook header is invalid (possible data tampering).",
        "unit": "VietQRInboundController\n.receiveCallback()",
        "input": "POST /vqr/bank/api/test/\ntransaction-callback\nAuthorization: Bearer INVALID\n(valid payload body)",
        "expected": "throws CallbackValidationException\nHTTP 400 Bad Request",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Alt flow:\nData Tampered",
    },
    {
        "id": "UT016",
        "name": "Callback – Payment Failed Status",
        "method": "shouldThrowPaymentFailedExceptionWhenPayloadNotSuccess",
        "description": "Throws PaymentFailedException when VietQR reports a failed transaction status (status=FAILED).",
        "unit": "VietQRInboundController\n.receiveCallback()",
        "input": "POST /vqr/bank/api/test/\ntransaction-callback\nAuthorization: Bearer dummy_token_123\n{\n  status: \"FAILED\",\n  amount: 100000\n}",
        "expected": "throws PaymentFailedException\nHTTP 400 Bad Request",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Alt flow:\nPayment Failed",
    },
    {
        "id": "UT017",
        "name": "Callback – Payment Session Not Found",
        "method": "shouldThrowPaymentFailedExceptionWhenSessionNotFound",
        "description": "Throws PaymentFailedException when no cached Invoice is found for the paymentRef extracted from the transfer content.",
        "unit": "VietQRInboundController\n.receiveCallback()",
        "input": "POST /vqr/bank/api/test/\ntransaction-callback\nAuthorization: Bearer dummy_token_123\n{\n  content: \"Payment AIMS REF-12345\"\n}\n(Cache has NO Invoice for REF-12345)",
        "expected": "throws PaymentFailedException\nHTTP 400 Bad Request",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "paymentRef not found in cache",
    },
    # ── PaymentCacheService ──
    {
        "id": "UT018",
        "name": "Cache – Add and Get Pending Invoice",
        "method": "shouldAddAndGetPendingInvoice",
        "description": "Store a temporary Invoice in the pending cache and retrieve it accurately by paymentRef.",
        "unit": "PaymentCacheService\n.addPendingInvoice()\n.getPendingInvoice()",
        "input": "paymentRef = \"REF-TEST-123\"\nInvoice { totalAmount: 100000 }",
        "expected": "getPendingInvoice(\"REF-TEST-123\")\n  -> Invoice is not null\n  -> totalAmount = 100000",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Verify add and retrieve",
    },
    {
        "id": "UT019",
        "name": "Cache – Remove Pending Invoice",
        "method": "shouldRemovePendingInvoice",
        "description": "Remove an Invoice from the pending cache by paymentRef; subsequent lookups should return null.",
        "unit": "PaymentCacheService\n.removePendingInvoice()\n.getPendingInvoice()",
        "input": "paymentRef = \"REF-REMOVE\"\n(Add Invoice first, then remove)",
        "expected": "getPendingInvoice(\"REF-REMOVE\")\n  -> null",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Verify removal from cache",
    },
    {
        "id": "UT020",
        "name": "Cache – Add and Get Completed Payment",
        "method": "shouldAddAndGetCompletedPayment",
        "description": "Store a completed TransactionInfo in the cache and retrieve it accurately by paymentRef.",
        "unit": "PaymentCacheService\n.addCompletedPayment()\n.getCompletedPayment()",
        "input": "paymentRef = \"REF-COMPLETE\"\nTransactionInfo {\n  status: \"SUCCESS\",\n  amount: 50000\n}",
        "expected": "getCompletedPayment(\"REF-COMPLETE\")\n  -> not null\n  -> status = \"SUCCESS\"\n  -> amount = 50000",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Verify completed payment cache",
    },
    {
        "id": "UT021",
        "name": "Cache – Return Null When Key Not Found",
        "method": "shouldReturnNullWhenKeyNotFound",
        "description": "Returns null when searching with a key that does not exist in either pending or completed cache.",
        "unit": "PaymentCacheService\n.getPendingInvoice()\n.getCompletedPayment()",
        "input": "paymentRef = \"NON-EXISTENT\"\n(No data added to cache)",
        "expected": "getPendingInvoice() -> null\ngetCompletedPayment() -> null",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Key does not exist in cache",
    },
    # ── VietQRController ──
    {
        "id": "UT022",
        "name": "Get Valid Access Token – Fetch New Token",
        "method": "shouldReturnNewTokenWhenCacheIsEmpty",
        "description": "Fetches a new access token from the VietQR API when the cache is empty or the token has expired.",
        "unit": "VietQRController\n.getValidAccessToken()",
        "input": "Token cache is empty\nVietQR API returns:\n{\"access_token\":\"token123\",\n \"expires_in\":3600}",
        "expected": "Returned token = \"token123\"\nboundary.getAccessToken() called exactly 1 time",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Cache miss – fetch new token",
    },
    {
        "id": "UT023",
        "name": "Get Valid Access Token – Empty Token Response",
        "method": "shouldThrowInvalidTokenExceptionWhenTokenIsEmpty",
        "description": "Throws InvalidTokenException when VietQR returns an empty access_token string.",
        "unit": "VietQRController\n.getValidAccessToken()",
        "input": "VietQR API returns:\n{\"access_token\":\"\",\n \"expires_in\":3600}",
        "expected": "throws InvalidTokenException",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Empty token returned from VietQR",
    },
    {
        "id": "UT024",
        "name": "Generate QR Code Successfully (VietQRController)",
        "method": "shouldGenerateQRCodeSuccessfully",
        "description": "VietQRController successfully generates a QR code: fetches token, calls QR generation API, and parses the result.",
        "unit": "VietQRController\n.generateQRCode()",
        "input": "Invoice {\n  paymentReference: \"REF-123\",\n  totalAmount: 100000\n}\nVietQR returns:\n{\"data\":{\"qrCode\":\"000201...\",\n \"qrDataURL\":\"data:image/...\"}}",
        "expected": "QRCode is not null\nqrCode.getQrCode() = \"000201...\"",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Happy path – QR generated successfully",
    },
    {
        "id": "UT025",
        "name": "Generate QR Code – Empty QR Response",
        "method": "shouldThrowQRCodeGenerationExceptionWhenQREmpty",
        "description": "Throws QRCodeGenerationException when VietQR returns an empty qrCode string in the response.",
        "unit": "VietQRController\n.generateQRCode()",
        "input": "Invoice { paymentReference: \"REF-123\" }\nVietQR returns:\n{\"data\":{\"qrCode\":\"\",\n \"qrDataURL\":\"\"}}",
        "expected": "throws QRCodeGenerationException",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "VietQR returned empty QR code",
    },
    {
        "id": "UT026",
        "name": "Check Payment Status – SUCCESS",
        "method": "shouldReturnSuccessWhenResponseIs00",
        "description": "Returns \"SUCCESS\" when the VietQR API responds with code \"00\" (transaction successful).",
        "unit": "VietQRController\n.checkPaymentStatus()",
        "input": "Invoice { paymentReference: \"REF-123\" }\nVietQR check returns:\n{\"code\":\"00\"}",
        "expected": "Result = \"SUCCESS\"",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Code \"00\" = success",
    },
    {
        "id": "UT027",
        "name": "Check Payment Status – FAILED",
        "method": "shouldReturnFailedWhenResponseIsFailed",
        "description": "Returns \"FAILED\" when the VietQR response contains the string \"FAILED\".",
        "unit": "VietQRController\n.checkPaymentStatus()",
        "input": "Invoice { paymentReference: \"REF-123\" }\nVietQR check returns:\n\"FAILED transaction\"",
        "expected": "Result = \"FAILED\"",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Response contains \"FAILED\"",
    },
    {
        "id": "UT028",
        "name": "Check Payment Status – PENDING",
        "method": "shouldReturnPendingWhenResponseIsUnknown",
        "description": "Returns \"PENDING\" when the VietQR response does not contain a success or failure indicator.",
        "unit": "VietQRController\n.checkPaymentStatus()",
        "input": "Invoice { paymentReference: \"REF-123\" }\nVietQR check returns:\n{\"code\":\"99\"}",
        "expected": "Result = \"PENDING\"",
        "actual": "TBD",
        "pass_fail": "TBD",
        "notes": "Code is neither \"00\" nor \"FAILED\"",
    },
]

# ─────────────────────────────────────────────
# Parse JUnit reports and map to TEST_CASES
# ─────────────────────────────────────────────
def parse_test_results():
    if not os.path.exists(REPORTS_DIR):
        print(f"⚠️ Reports directory not found at {REPORTS_DIR}. Using default 'TBD' values.")
        return

    results = {}  # maps method_name -> (status, time)
    for file_name in os.listdir(REPORTS_DIR):
        if file_name.startswith("TEST-") and file_name.endswith(".xml"):
            file_path = os.path.join(REPORTS_DIR, file_name)
            try:
                tree = ET.parse(file_path)
                root = tree.getroot()
                for tc in root.findall("testcase"):
                    name = tc.get("name")
                    time_val = tc.get("time", "0.0")
                    
                    # Check for failures or errors
                    failure = tc.find("failure")
                    error = tc.find("error")
                    
                    if failure is not None:
                        msg = failure.get("message", "Test Failed")
                        results[name] = (f"FAILED ({msg})", "FAIL")
                    elif error is not None:
                        msg = error.get("message", "Error occurred")
                        results[name] = (f"ERROR ({msg})", "FAIL")
                    else:
                        results[name] = (f"PASSED (Duration: {float(time_val):.3f}s)", "PASS")
            except Exception as e:
                print(f"Error parsing report {file_name}: {e}")

    # Apply to our test cases list
    for tc in TEST_CASES:
        method = tc.get("method")
        if method in results:
            actual, pass_fail = results[method]
            tc["actual"] = actual
            tc["pass_fail"] = pass_fail
        else:
            tc["actual"] = "TBD"
            tc["pass_fail"] = "TBD"

# ─────────────────────────────────────────────
# Docx Generation Utilities
# ─────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_cell_text(cell, text, bold=False, font_size=9, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = 'Calibri'
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def create_document():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2)
        section.right_margin  = Cm(2)

    # Main title
    title = doc.add_heading("5. Test Case Details", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.runs[0]
    title_run.font.size = Pt(14)
    title_run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    intro = doc.add_paragraph(
        "Use Case: Pay Order (Order Payment via VietQR Payment Gateway)\n"
        "Project: AIMS Payment Service\n"
        "Test cases aggregated from: PayOrderControllerTest, "
        "VietQRInboundControllerTest, PaymentCacheServiceTest, VietQRControllerTest."
    )
    intro.runs[0].font.size = Pt(10)

    doc.add_paragraph()

    # Table headers
    COLS = [
        "Test Case ID", "Test Case Name", "Description",
        "Unit Under Test", "Input Data",
        "Expected Output", "Actual Output", "Pass/Fail", "Notes"
    ]
    COL_WIDTHS = [
        Cm(1.6), Cm(3.5), Cm(4.5), Cm(3.0), Cm(4.8),
        Cm(4.5), Cm(2.5), Cm(1.8), Cm(2.5)
    ]

    table = doc.add_table(rows=1, cols=len(COLS))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    HEADER_COLOR = "1F497D"
    hdr_cells = table.rows[0].cells
    for i, (col_name, width) in enumerate(zip(COLS, COL_WIDTHS)):
        hdr_cells[i].width = width
        set_cell_bg(hdr_cells[i], HEADER_COLOR)
        add_cell_text(hdr_cells[i], col_name, bold=True, font_size=9,
                      color=(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)

    # Data rows
    EVEN_ROW_COLOR = "EBF3FB"
    ODD_ROW_COLOR  = "FFFFFF"

    for idx, tc in enumerate(TEST_CASES):
        row = table.add_row()
        bg = EVEN_ROW_COLOR if idx % 2 == 0 else ODD_ROW_COLOR

        values = [
            tc["id"],
            tc["name"],
            tc["description"],
            tc["unit"],
            tc["input"],
            tc["expected"],
            tc["actual"],
            tc["pass_fail"],
            tc["notes"],
        ]

        for i, (val, width) in enumerate(zip(values, COL_WIDTHS)):
            cell = row.cells[i]
            cell.width = width
            set_cell_bg(cell, bg)
            
            # Make PASS green, FAIL red
            font_color = None
            if i == 7: # Pass/Fail column
                if val == "PASS":
                    font_color = (0x38, 0x76, 0x1D) # Dark green
                elif val == "FAIL":
                    font_color = (0xCC, 0x00, 0x00) # Dark red

            add_cell_text(
                cell, val, font_size=8.5, color=font_color,
                align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 7) else WD_ALIGN_PARAGRAPH.LEFT
            )

    doc.save(OUTPUT_PATH)
    print(f"[OK] Saved: {OUTPUT_PATH}")

if __name__ == "__main__":
    parse_test_results()
    create_document()
