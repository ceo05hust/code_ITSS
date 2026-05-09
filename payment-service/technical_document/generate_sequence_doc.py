import os
import sys
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    os.system(f"{sys.executable} -m pip install python-docx")
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

doc = Document()

# Title
title = doc.add_heading('Tài Liệu Kỹ Thuật: Luồng Hoạt Động (Sequence) Payment Service', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Tài liệu này mô tả chi tiết luồng gọi API, cách xử lý logic bên trong code, và cách Webhook hoạt động trong dự án tích hợp thanh toán VietQR.')

# 1. Tổng quan Kiến trúc
doc.add_heading('1. Tổng quan các Thành phần (Components)', level=1)
p1 = doc.add_paragraph()
p1.add_run('Hệ thống chia làm 3 thành phần chính giao tiếp với nhau:\n')
p1.add_run('1. Frontend (Angular): ').bold = True
p1.add_run('Cung cấp giao diện tương tác cho người dùng (Xem đơn hàng, Quét QR, Bấm xác nhận).\n')
p1.add_run('2. Backend (Spring Boot): ').bold = True
p1.add_run('Đóng vai trò "nhạc trưởng". Gồm các Controller phục vụ Frontend (PayOrderController) và các Controller phục vụ Webhook (VietQRInboundController).\n')
p1.add_run('3. Hệ thống VietQR (External API): ').bold = True
p1.add_run('Ngân hàng mẹ. Cung cấp API tạo QR và chủ động đẩy (Push) thông báo khi nhận được tiền.')

# 2. Luồng Tạo mã QR
doc.add_heading('2. Luồng 1: Tạo mã QR (Generate QR Flow)', level=1)
doc.add_paragraph('Xảy ra khi người dùng bấm nút "Thanh toán bằng VietQR" trên giao diện.')
p2 = doc.add_paragraph()
p2.add_run('Bước 1: ').bold = True
p2.add_run('Frontend gọi API ')
p2.add_run('POST /api/payment/generate-qr').italic = True
p2.add_run(' truyền vào thông tin đơn hàng (InvoiceId, TotalAmount).\n')
p2.add_run('Bước 2: ').bold = True
p2.add_run('Tại Backend, PayOrderController nhận request, gọi xuống lớp VietQRController.\n')
p2.add_run('Bước 3: ').bold = True
p2.add_run('VietQRController đầu tiên sẽ nhờ VietQRBoundary gọi API lấy Token của VietQR (bằng ClientId/Secret hoặc Username/Password).\n')
p2.add_run('Bước 4: ').bold = True
p2.add_run('Có Token, VietQRController đóng gói Payload (bankCode, amount, content...) và gọi tiếp API ')
p2.add_run('POST /vqr/api/qr/generate-customer').italic = True
p2.add_run(' của VietQR.\n')
p2.add_run('Bước 5: ').bold = True
p2.add_run('VietQR trả về chuỗi JSON chứa qrCode (dạng text) và qrLink (URL ảnh). Backend bóc tách và trả thẳng về cho Frontend hiển thị ảnh QR lên màn hình.')

# 3. Luồng Webhook
doc.add_heading('3. Luồng 2: Nhận thông báo tự động (Webhook / IPN Flow)', level=1)
doc.add_paragraph('Đây là luồng "Bị động" (Inbound). Xảy ra khi người dùng cầm điện thoại quét mã và chuyển tiền thành công. VietQR sẽ chủ động "gõ cửa" hệ thống của chúng ta.')
p3 = doc.add_paragraph()
p3.add_run('Bước 1 (Xác thực 2 chiều): ').bold = True
p3.add_run('VietQR gọi API ')
p3.add_run('POST /vqr/api/token_generate').italic = True
p3.add_run(' của chúng ta, kèm theo Basic Auth (admin:admin). VietQRInboundController kiểm tra mật khẩu và trả về một chiếc chìa khoá là "dummy_token_123".\n')
p3.add_run('Bước 2 (Nhận dữ liệu): ').bold = True
p3.add_run('VietQR cầm chìa khoá đó, gọi tiếp vào API Webhook của ta: ')
p3.add_run('POST /vqr/bank/api/test/transaction-callback').italic = True
p3.add_run(' kèm theo thông tin giao dịch (RefId, Số tiền, Content).\n')
p3.add_run('Bước 3 (Xử lý dữ liệu): ').bold = True
p3.add_run('VietQRInboundController kiểm tra token. Nếu hợp lệ, nó sẽ: (1) Lưu bản ghi vào bảng TransactionInfo, (2) Tìm hoá đơn (Invoice) tương ứng, (3) Đánh dấu hoá đơn thành PAID (Đã thanh toán) và lưu xuống Database.')

# 4. Luồng Confirm
doc.add_heading('4. Luồng 3: Xác nhận thanh toán (Confirm Payment Flow)', level=1)
doc.add_paragraph('Luồng này xảy ra khi khách hàng bấm nút "Tôi đã chuyển tiền xong" (hoặc Frontend tự động polling liên tục để hỏi Backend xem đơn hàng đã xong chưa).')
p4 = doc.add_paragraph()
p4.add_run('Bước 1: ').bold = True
p4.add_run('Frontend gọi API ')
p4.add_run('POST /api/payment/confirm').italic = True
p4.add_run(' truyền vào invoiceId.\n')
p4.add_run('Bước 2 (Kiểm tra chéo): ').bold = True
p4.add_run('PayOrderController tìm Invoice trong Database. Nó kiểm tra xem Invoice đã có TransactionInfo gắn vào chưa (nghĩa là Webhook ở Luồng 2 đã chạy xong chưa).\n')
p4.add_run('Trường hợp 2A (Thành công): ').bold = True
p4.add_run('Nếu Invoice đã được đánh dấu PAID (do Webhook đã kịp xử lý), Backend lập tức trả về 200 OK "Invoice already paid". Frontend nhảy sang trang Thành Công (Tick xanh).\n')
p4.add_run('Trường hợp 2B (Sandbox Simulation): ').bold = True
p4.add_run('Nếu Invoice chưa PAID, Backend sẽ thử ép VietQR giả lập chuyển tiền bằng cách gọi API checkPaymentStatus sang VietQR (chỉ dùng cho môi trường Test). Nếu VietQR đồng ý giả lập, vòng lặp Webhook sẽ được kích hoạt.')

doc.save('Luong_Ky_Thuat_Payment_Service.docx')
print("Document saved as Luong_Ky_Thuat_Payment_Service.docx")
