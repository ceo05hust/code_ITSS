import os
import sys
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    os.system(f"{sys.executable} -m pip install python-docx")
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('Tài Liệu Ghi Nhớ: Vấn đề Hardcode Invoice ID (INV-001)', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Tài liệu này ghi chú lại việc gán cứng (hardcode) dữ liệu mã đơn hàng trong phiên bản Frontend hiện tại. Mục đích để nhóm phát triển nhớ thay đổi lại thành dữ liệu động (Dynamic Data) khi tích hợp hoàn chỉnh với hệ thống Giỏ hàng (Cart) của AIMS.')

# 1. Tình trạng hiện tại
doc.add_heading('1. Tình trạng hiện tại (Tại sao lại hardcode?)', level=1)
doc.add_paragraph('Hiện tại, hệ thống Frontend của Payment Service đang đóng vai trò giả lập độc lập. Vì chưa có hệ thống Quản lý Đơn hàng (Order Management) hay Giỏ hàng (Cart) đứng trước nó để sinh ra mã hoá đơn thật, chúng ta phải truyền cố định chuỗi "INV-001" vào các API để Backend có dữ liệu kiểm thử luồng sinh mã QR và luồng xác nhận thanh toán.')

# 2. Vị trí các dòng code cần sửa
doc.add_heading('2. Các vị trí Code đang bị Hardcode cần sửa lại sau này', level=1)
doc.add_paragraph('Bạn cần lưu ý sửa lại các vị trí sau trong mã nguồn Frontend khi ghép nối hệ thống:')

doc.add_heading('A. File: frontend/src/app/components/qr-scan/qr-scan.component.ts', level=2)
p2a1 = doc.add_paragraph()
p2a1.add_run('Hàm generateQR()').bold = True
p2a1.add_run(': Đang gán cứng object gửi xuống API.\nCode hiện tại:\n')
p2a1.style = 'Normal'
code1 = doc.add_paragraph()
code1.add_run("this.paymentService.generateQr({\n  invoiceId: 'INV-001', // <--- CẦN SỬA\n  shippingFee: 30000,\n  totalProductPriceExVAT: 200000,\n  totalProductPriceIncVAT: 220000,\n  totalAmount: 250000\n}).subscribe({ ... });")
code1.style = 'Intense Quote'

p2a2 = doc.add_paragraph()
p2a2.add_run('Hàm confirmPayment()').bold = True
p2a2.add_run(': Đang gọi cứng mã để kiểm tra thanh toán.\nCode hiện tại:\n')
code2 = doc.add_paragraph()
code2.add_run("this.paymentService.confirmPayment({ invoiceId: 'INV-001' }).subscribe({ ... });")
code2.style = 'Intense Quote'

doc.add_heading('B. Giao diện HTML (Hiển thị)', level=2)
p2b = doc.add_paragraph()
p2b.add_run('Trong các file ')
p2b.add_run('checkout.component.ts').italic = True
p2b.add_run(' và ')
p2b.add_run('success.component.ts').italic = True
p2b.add_run(', giao diện đang in trực tiếp text ra màn hình.\n')
code3 = doc.add_paragraph()
code3.add_run("<strong>INV-001</strong>")
code3.style = 'Intense Quote'

# 3. Hướng giải quyết thực tế (Production)
doc.add_heading('3. Hướng giải quyết khi làm thật (Production)', level=1)
p3 = doc.add_paragraph()
p3.add_run('Khi ghép nối thành dự án AIMS hoàn chỉnh, luồng chuẩn sẽ diễn ra như sau:\n')
p3.add_run('1. Bước chọn món: ').bold = True
p3.add_run('Khách hàng bấm chọn đĩa CD/Sách, điền thông tin giao hàng rồi bấm nút "Xác nhận đặt hàng".\n')
p3.add_run('2. Backend tạo Invoice: ').bold = True
p3.add_run('Hệ thống AIMS (OrderController) sẽ lưu đơn hàng vào Database và sinh ra một mã Invoice ngẫu nhiên, ví dụ: ')
p3.add_run('INV-823901').italic = True
p3.add_run('.\n')
p3.add_run('3. Chuyển hướng (Routing): ').bold = True
p3.add_run('AIMS sẽ điều hướng người dùng sang trang Thanh toán của Payment Service (Giao diện Frontend chúng ta vừa làm), kèm theo mã Invoice thông qua URL Parameter hoặc State.\nVí dụ URL: ')
p3.add_run('http://localhost:4200/qr-scan?invoiceId=INV-823901').italic = True
p3.add_run('.\n')
p3.add_run('4. Lấy dữ liệu động: ').bold = True
p3.add_run('Lúc này, trong hàm ')
p3.add_run('ngOnInit()').italic = True
p3.add_run(' của mã nguồn Angular, thay vì dùng chuỗi tĩnh, ta sẽ đọc mã hoá đơn từ thanh URL:\n')

code4 = doc.add_paragraph()
code4.add_run("this.route.queryParams.subscribe(params => {\n  this.currentInvoiceId = params['invoiceId'];\n  // Gọi API với mã động\n  this.paymentService.generateQr({ invoiceId: this.currentInvoiceId, ... });\n});")
code4.style = 'Intense Quote'

doc.save('GhiNho_Hardcode_InvoiceId.docx')
print("Document saved as GhiNho_Hardcode_InvoiceId.docx")
