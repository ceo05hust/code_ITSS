import os
import sys
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    os.system(f"{sys.executable} -m pip install python-docx")
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('Tổng Kết Kiến Thức Project Payment Service (VietQR)', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Intro
doc.add_paragraph('Tài liệu này tổng hợp lại các vấn đề đã gặp phải, cách giải quyết và những bài học rút ra trong phiên làm việc ngày hôm nay.')

# Section 1
doc.add_heading('1. Vấn đề với Webhook và Lỗi E76 (VietQR)', level=1)
p1 = doc.add_paragraph()
p1.add_run('Tình huống: ').bold = True
p1.add_run('Khi chạy API Confirm, hệ thống liên tục báo lỗi 400 Bad Request kèm mã lỗi nội bộ E76 từ VietQR.\n')
p1.add_run('Nguyên nhân: ').bold = True
p1.add_run('Mã E76 có nghĩa là "Merchant chưa được cấu hình / chưa đăng ký hệ thống". Điều này xảy ra do Webhook URL (ngrok) chưa được điền vào portal Sandbox của VietQR, hoặc tài khoản đăng ký bị sai mật khẩu (do copy thiếu/nhầm mã Base64).\n')
p1.add_run('Bài học lớn nhất: ').bold = True
p1.add_run(' Ngay cả khi đã điền thành công Webhook trên Portal Sandbox, VietQR vẫn yêu cầu ')
p1.add_run('Admin VietQR phải phê duyệt thủ công').underline = True
p1.add_run(' thì webhook mới bắt đầu hoạt động. Trong lúc chờ duyệt, gọi API test callback vẫn sẽ ném lỗi E76.')

# Section 2
doc.add_heading('2. Giải pháp giả lập (Bypass) Webhook', level=1)
p2 = doc.add_paragraph()
p2.add_run('Vì phải chờ duyệt từ phía VietQR, chúng ta đã dùng thủ thuật giả lập (mocking) để tiếp tục luồng code:\n')
p2.add_run('Cách làm: ').bold = True
p2.add_run('Thay vì chờ VietQR gọi về hệ thống, chúng ta tự dùng PowerShell để bắn trực tiếp một HTTP POST request vào đường dẫn ')
p2.add_run('http://localhost:8080/vqr/bank/api/test/transaction-callback').italic = True
p2.add_run(' kèm theo "Bearer dummy_token_123" và nội dung giao dịch.\n')
p2.add_run('Ý nghĩa: ').bold = True
p2.add_run('Giúp hệ thống AIMS lách qua bước chờ đợi của VietQR, tự động cập nhật Database thành trạng thái PAID (Đã thanh toán) để có thể test thẳng API Confirm.')

# Section 3
doc.add_heading('3. Phân biệt Backend và Frontend', level=1)
p3 = doc.add_paragraph()
p3.add_run('Câu hỏi: ').bold = True
p3.add_run('"Đây mới là backend thôi phải không?"\n')
p3.add_run('Giải đáp: ').bold = True
p3.add_run('Chính xác. Toàn bộ code Java (Spring Boot) chạy trên cổng 8080 chỉ xử lý logic ngầm, kết nối CSDL và giao tiếp với ngân hàng. Để hệ thống hoàn chỉnh, chúng ta cần một Frontend (Giao diện) để người dùng thực sự thao tác.')

# Section 4
doc.add_heading('4. Tích hợp Frontend (Angular 18)', level=1)
p4 = doc.add_paragraph()
p4.add_run('Chúng ta đã tự tay xây dựng một giao diện Frontend bằng Angular và TypeScript trong thư mục payment-service/frontend.\n')
p4.add_run('- Thiết kế: ').bold = True
p4.add_run('Sử dụng phong cách Glassmorphism (Kính mờ) hiện đại, có các màn hình Checkout, Quét mã QR và Thông báo thành công.\n')
p4.add_run('- Kết nối: ').bold = True
p4.add_run('Sử dụng file proxy.conf.json để giải quyết lỗi CORS khi Frontend (cổng 4200) gọi xuống Backend (cổng 8080).')

# Section 5
doc.add_heading('5. Lời nhắc (Warning) về Unused Import của IDE', level=1)
p5 = doc.add_paragraph()
p5.add_run('Câu hỏi: ').bold = True
p5.add_run('"Sao dòng import QRAccessTokenResponse lại bị gạch đi vậy?"\n')
p5.add_run('Giải đáp: ').bold = True
p5.add_run('Đó là tính năng "Unused Import" của IDE (trình soạn thảo). Vì trong quá trình tối ưu code (refactor) hàm getAccessToken, chúng ta đã chuyển sang trả về chuỗi String thô (raw JSON) nên class QRAccessTokenResponse không còn được dùng nữa. Việc xoá bỏ các dòng import thừa này giúp code gọn gàng hơn, không ảnh hưởng gì tới luồng chạy.')

# Section 6
doc.add_heading('6. Vấn đề bảo mật khi đẩy code lên GitHub', level=1)
p6 = doc.add_paragraph()
p6.add_run('Câu hỏi: ').bold = True
p6.add_run('"Đẩy code lên Github có nguy hiểm không?"\n')
p6.add_run('Giải đáp: ').bold = True
p6.add_run('Do repo đang để chế độ Private và tài khoản VietQR đang dùng là tài khoản Sandbox (thử nghiệm) nên hoàn toàn an toàn.\n')
p6.add_run('Lưu ý (Best Practice): ').bold = True
p6.add_run('Trong môi trường thực tế (Production), tuyệt đối không được push các file chứa mật khẩu (như application.properties hay các file test.ps1) lên Github. Thường lập trình viên sẽ dùng file .env và đưa nó vào .gitignore.')

# Section 7
doc.add_heading('7. Cẩm nang Test Backend bằng PowerShell (Khi chưa có Frontend)', level=1)
doc.add_paragraph('Để giả lập luồng thanh toán từ A-Z chỉ bằng Terminal (PowerShell) mà không cần giao diện, hãy chạy các lệnh sau theo thứ tự:')

doc.add_heading('Bước 1: Tạo mã QR (Generate QR)', level=2)
doc.add_paragraph("Lệnh này giả lập lúc user nhấn nút Thanh toán để lấy mã QR về:")
cmd1 = doc.add_paragraph()
cmd1.add_run("Invoke-RestMethod -Uri \"http://localhost:8080/api/payment/generate-qr\" -Method POST -ContentType \"application/json\" -Body '{\"invoiceId\":\"INV-001\",\"shippingFee\":30000,\"totalProductPriceExVAT\":200000,\"totalProductPriceIncVAT\":220000,\"totalAmount\":250000}'")
cmd1.style = 'Intense Quote'

doc.add_heading('Bước 2: Giả lập Webhook (Khi VietQR chưa duyệt)', level=2)
doc.add_paragraph("Lệnh này bắn thẳng vào localhost của bạn để báo cáo rằng khách hàng đã chuyển tiền (vượt qua việc VietQR bị lỗi E76):")
cmd2 = doc.add_paragraph()
cmd2.add_run("Invoke-RestMethod -Uri \"http://localhost:8080/vqr/bank/api/test/transaction-callback\" -Method POST -Headers @{ Authorization = \"Bearer dummy_token_123\" } -ContentType \"application/json\" -Body '{\"transactionRefId\":\"INV-001\", \"transactionId\":\"VQR-9999\", \"amount\":250000, \"status\":\"00\", \"success\":true, \"message\":\"Success\", \"content\":\"AIMS INV-001\"}'")
cmd2.style = 'Intense Quote'

doc.add_heading('Bước 3: Xác nhận đơn hàng (Confirm)', level=2)
doc.add_paragraph("Lệnh này kiểm tra xem đơn hàng đã được cập nhật thành PAID hay chưa:")
cmd3 = doc.add_paragraph()
cmd3.add_run("Invoke-RestMethod -Uri \"http://localhost:8080/api/payment/confirm\" -Method POST -ContentType \"application/json\" -Body '{\"invoiceId\":\"INV-001\"}'")
cmd3.style = 'Intense Quote'

# Section 8
doc.add_heading('8. Hướng dẫn khởi động hệ thống Hoàn Chỉnh (Backend + Frontend)', level=1)
doc.add_paragraph('Khi đã có đầy đủ Frontend và Backend, bạn chỉ cần mở 2 cửa sổ PowerShell riêng biệt để bật cả 2 lên:')

doc.add_heading('Terminal 1: Khởi động Backend (Spring Boot)', level=2)
doc.add_paragraph("Mở thư mục payment-service và chạy:")
cmd4 = doc.add_paragraph()
cmd4.add_run(".\\mvnw.cmd spring-boot:run")
cmd4.style = 'Intense Quote'

doc.add_heading('Terminal 2: Khởi động Frontend (Angular)', level=2)
doc.add_paragraph("Mở thư mục payment-service/frontend và chạy:")
cmd5 = doc.add_paragraph()
cmd5.add_run("npm start")
cmd5.style = 'Intense Quote'
doc.add_paragraph("Sau khi Terminal báo Compiled successfully, bạn mở trình duyệt và truy cập: http://localhost:4200")

doc.save('Tong_Ket_Project_VietQR_v2.docx')
print("Document saved as Tong_Ket_Project_VietQR_v2.docx")
